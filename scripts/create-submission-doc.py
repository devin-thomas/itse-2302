from pathlib import Path
import json

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE_DIR = ROOT / "output" / "playwright" / "windows-lab-1"
OUTPUT = EVIDENCE_DIR / "ITSE-2302-Lab-1-Windows-Submission.docx"


def set_font(style, name, size, color, bold=False, italic=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    set_font(normal, "Calibri", 11, "172B4D")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = document.styles["Title"]
    set_font(title, "Calibri", 24, "0B2545", bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)

    heading = document.styles["Heading 1"]
    set_font(heading, "Calibri", 16, "2E74B5", bold=True)
    heading.paragraph_format.space_before = Pt(18)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True

    if "Evidence Caption" not in [style.name for style in document.styles]:
        caption = document.styles.add_style("Evidence Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = document.styles["Evidence Caption"]
    set_font(caption, "Calibri", 9, "5B677A", italic=True)
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    footer_run = footer.add_run("ITSE 2302 | Lab 1 | Windows evidence")
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("7A869A")

    document.core_properties.title = "ITSE 2302 Lab 1 Windows Submission"
    document.core_properties.subject = "XAMPP installation and test evidence"
    document.core_properties.author = "ITSE 2302"
    document.core_properties.last_modified_by = "ITSE 2302"


def add_evidence(document, number, title, description, filename, width):
    if number > 1:
        document.add_page_break()

    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.keep_with_next = True
    heading.add_run(f"{number}. {title}")

    detail = document.add_paragraph()
    detail.paragraph_format.space_after = Pt(8)
    detail_run = detail.add_run(description)
    detail_run.font.name = "Calibri"
    detail_run.font.size = Pt(10.5)
    detail_run.font.color.rgb = RGBColor.from_string("526173")

    image_path = EVIDENCE_DIR / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(3)
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(width))

    caption = document.add_paragraph(style="Evidence Caption")
    caption.add_run(f"Figure {number}. {description}")


def main():
    report_path = EVIDENCE_DIR / "evidence.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    mysql_port = report.get("mysqlPort", 3306)

    document = Document()
    configure_document(document)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(3)
    kicker_run = kicker.add_run("ITSE 2302 | INTERMEDIATE WEB")
    kicker_run.font.name = "Calibri"
    kicker_run.font.size = Pt(9)
    kicker_run.font.bold = True
    kicker_run.font.color.rgb = RGBColor.from_string("2E74B5")

    title = document.add_paragraph(style="Title")
    title.add_run("Lab 1 Submission")

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run("XAMPP Installation and Test - Windows")
    subtitle_run.font.name = "Calibri"
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor.from_string("526173")

    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(12)
    metadata_run = metadata.add_run("Student: Devin Thomas    |    Platform: Windows    |    Date: August 30, 2026")
    metadata_run.font.name = "Calibri"
    metadata_run.font.size = Pt(10.5)
    metadata_run.font.bold = True
    metadata_run.font.color.rgb = RGBColor.from_string("172B4D")

    note = document.add_paragraph()
    note.paragraph_format.space_after = Pt(12)
    note_run = note.add_run(
        f"The screenshots below document the required local XAMPP setup. Browser captures use a clean headed Chromium profile, include the address bar, and verify Apache on port 80 and XAMPP MariaDB on port {mysql_port}."
    )
    note_run.font.name = "Calibri"
    note_run.font.size = Pt(10.5)
    note_run.font.color.rgb = RGBColor.from_string("526173")

    add_evidence(
        document,
        1,
        "XAMPP Control Panel",
        "Apache and MySQL/MariaDB are started in the Windows XAMPP Control Panel.",
        "control-panel-window.png",
        6.0,
    )
    add_evidence(
        document,
        2,
        "XAMPP Welcome Page",
        "The browser address bar shows localhost/dashboard/ and the XAMPP welcome page is visible.",
        "welcome-window.png",
        6.5,
    )
    add_evidence(
        document,
        3,
        "Project Page",
        "The browser address bar shows localhost/project/index.php and the required name is visible.",
        "project-window.png",
        6.5,
    )
    add_evidence(
        document,
        4,
        "Live Status Check",
        f"The local PHP status page confirms Apache and MySQL/MariaDB are RUNNING on the selected port {mysql_port}.",
        "status-window.png",
        6.5,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
